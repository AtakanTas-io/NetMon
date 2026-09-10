import json
import time

import pytest
import server


@pytest.fixture()
def isolated_application(tmp_path, monkeypatch):
    monkeypatch.setattr(server, "DB_PATH", tmp_path / "application.db")
    monkeypatch.setattr(server, "INITIAL_PASSWORD_PATH", tmp_path / "initial-admin.txt")
    server._devices_cache.update(
        {
            "ts": 0,
            "data": [],
            "error": None,
            "scan_status": "idle",
            "active_networks": [],
        }
    )
    server.init_db()
    return tmp_path


def _inventory_device(ip, mac, status="online", ports=None):
    return {
        "ip": ip,
        "mac": mac,
        "hostname": f"host-{ip.rsplit('.', 1)[-1]}",
        "vendor": "Test Vendor",
        "type": "computer",
        "status": status,
        "last_seen": time.time(),
        "network_id": 1,
        "network_cidr": "192.168.10.0/24",
        "discovery_sources": ["arp"],
        "connectivity_status": "online" if status == "online" else status,
        "identification_status": "identified",
        "classification": {"confidence": 0.8, "open_ports": ports or []},
        "wmi_inventory": {
            "hardware": {"cpu_model": "Test CPU", "ram_gb": 16},
            "software": {"os_name": "Test OS"},
            "security": {"antivirus": "Test AV", "firewall": "Açık"},
            "storage": [{"drive_letter": "C", "total_gb": 512}],
        },
    }


def test_legacy_successful_inventory_requires_detailed_refresh():
    legacy = {"status": "Success", "hardware": {"cpu_model": "Eski CPU"}}
    current = {"status": "Success", "inventory_schema_version": server.INVENTORY_PAYLOAD_SCHEMA_VERSION}

    assert server._inventory_payload_needs_upgrade(legacy) is True
    assert server._inventory_payload_needs_upgrade(current) is False


def test_enrich_devices_persists_new_identity_and_restores_manual_fields(isolated_application, monkeypatch):
    monkeypatch.setattr(
        server.diag,
        "get_network_context",
        lambda: {
            "local_ip": "192.168.10.5",
            "local_mac": "AA-BB-CC-DD-EE-05",
            "cidr": "192.168.10.0/24",
        },
    )
    first = _inventory_device("192.168.10.5", None, ports=[22])
    first["is_self"] = True
    created = server.enrich_devices([first])[0]

    assert created["mac"] == "AA:BB:CC:DD:EE:05"
    assert created["is_new"] is True
    with server.db_conn() as conn:
        conn.execute(
            "UPDATE known_devices SET friendly_name=?,owner=?,notes=?,device_type=?,classification_source=?,open_ports=? WHERE mac=?",
            ("Elle adlandırıldı", "BT", "Korunacak not", "server", "manual", json.dumps([22]), created["mac"]),
        )
        conn.commit()

    changed = _inventory_device("192.168.10.6", created["mac"], ports=[22, 443])
    changed["hostname"] = None
    restored = server.enrich_devices([changed])[0]

    assert restored["friendly_name"] == "Elle adlandırıldı"
    assert restored["owner"] == "BT"
    assert restored["type"] == "server"
    with server.db_conn() as conn:
        assert conn.execute("SELECT COUNT(*) FROM alerts WHERE message LIKE '%YENI PORT%'").fetchone()[0] == 1


def test_enrich_devices_keeps_unseen_identity_offline_and_handles_macless_device(isolated_application, monkeypatch):
    monkeypatch.setattr(server.diag, "get_network_context", lambda: {"cidr": "10.0.0.0/24"})
    with server.db_conn() as conn:
        conn.execute(
            "INSERT INTO known_devices(mac,friendly_name,hostname,device_type,first_seen,last_seen,last_ip,last_vendor,last_discovery_sources,identification_status,last_network) "
            "VALUES(?,?,?,?,?,?,?,?,?,?,?)",
            (
                "00:11:22:33:44:55",
                "Eski cihaz",
                "old-host",
                "printer",
                time.time() - 60,
                time.time() - 30,
                "10.0.0.9",
                "Printer Inc",
                json.dumps(["arp"]),
                "identified",
                "10.0.0.0/24",
            ),
        )
        conn.commit()

    result = server.enrich_devices(
        [
            {
                "ip": "10.0.0.8",
                "hostname": "macless",
                "status": "discovered",
                "friendly_name": "MAC bekleniyor",
                "classification": {},
            }
        ]
    )

    current = next(item for item in result if item["ip"] == "10.0.0.8")
    previous = next(item for item in result if item["ip"] == "10.0.0.9")
    assert current["mac"] == ""
    assert previous["status"] == "offline"
    assert previous["type"] == "printer"


def test_csv_export_covers_inventory_shapes_and_statuses(isolated_application):
    devices = [
        _inventory_device("192.168.10.2", "00:00:00:00:00:02", "online", [22, 443]),
        _inventory_device("192.168.10.3", "00:00:00:00:00:03", "offline"),
        _inventory_device("192.168.10.4", "00:00:00:00:00:04", "discovered"),
        _inventory_device("192.168.10.5", "00:00:00:00:00:05", "unknown"),
    ]
    devices[1]["classification"] = {}
    devices[1]["open_ports"] = "not-json"
    server._devices_cache["data"] = devices

    response = server.export_devices_csv()
    content = response.body.decode("utf-8-sig")

    assert response.status_code == 200
    assert "Çevrimiçi" in content
    assert "Çevrimdışı" in content
    assert "Yanıt Doğrulanamadı" in content
    assert "Belirsiz" in content
    assert "22, 443" in content


def test_csv_export_loads_known_devices_when_cache_is_empty(isolated_application):
    with server.db_conn() as conn:
        conn.execute(
            "INSERT INTO known_devices(mac,hostname,device_type,first_seen,last_seen,last_ip,last_status,open_ports) VALUES(?,?,?,?,?,?,?,?)",
            ("00:11:22:33:44:66", "db-host", "server", time.time(), time.time(), "10.0.0.6", "online", "[80]"),
        )
        conn.commit()

    response = server.export_devices_csv()

    assert response.status_code == 200
    assert "db-host" in response.body.decode("utf-8-sig")


def test_direct_xlsx_save_uses_available_folder(isolated_application, monkeypatch, tmp_path):
    from openpyxl import load_workbook

    server._devices_cache["data"] = [_inventory_device("10.0.0.2", "00:11:22:33:44:77", ports=[80])]
    monkeypatch.setattr(server.os.path, "expanduser", lambda value: str(tmp_path))
    monkeypatch.setattr(server.platform, "system", lambda: "Linux")
    with server.db_conn() as conn:
        now = time.time()
        conn.execute(
            """INSERT INTO connections(username,process_name,local_ip,remote_ip,remote_port,protocol,
            direction,status,resolved_hostname,first_seen,last_seen) VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
            (
                "DOMAIN\\ayse",
                "browser.exe",
                "10.0.0.2",
                "203.0.113.5",
                443,
                "TCP",
                "outbound",
                "ESTABLISHED",
                "example.test",
                now,
                now,
            ),
        )

    result = server.export_devices_save_to_disk(user={"username": "admin", "role": "admin"})

    assert result["ok"] is True
    assert result["count"] == 1
    assert result["filename"].endswith(".xlsx")
    saved = tmp_path / result["filename"]
    assert saved.exists()
    workbook = load_workbook(saved)
    assert {"Özet", "Cihaz Envanteri", "Donanım ve Yazılım", "Bağlantı Geçmişi", "Rapor Bilgisi"}.issubset(
        workbook.sheetnames
    )
    assert workbook.sheetnames[0] == "Özet"
    assert workbook["Cihaz Envanteri"]["A5"].value == "10.0.0.2"
    assert workbook["Bağlantı Geçmişi"]["A5"].value == "DOMAIN\\ayse"
    assert workbook["Bağlantı Geçmişi"]["B5"].value == "browser.exe"
    assert workbook["Cihaz Envanteri"].freeze_panes == "A5"
    assert workbook["Cihaz Envanteri"].sheet_view.showGridLines is False
    assert result["connection_count"] == 1


def test_xlsx_omits_connection_history_without_permission(isolated_application):
    from io import BytesIO

    from openpyxl import load_workbook

    content, connection_count = server._build_devices_xlsx([], include_connections=False)
    workbook = load_workbook(BytesIO(content), read_only=True)

    assert "Bağlantı Geçmişi" not in workbook.sheetnames
    assert connection_count == 0


@pytest.mark.parametrize(
    ("scenario", "expected_status"),
    [
        ("outage", "fail"),
        ("high_latency", "warn"),
        ("packet_loss", "warn"),
        ("dns_failure", "fail"),
        ("traffic_spike", "ok"),
        ("unknown", "unknown"),
    ],
)
def test_simulated_snapshots_cover_every_declared_scenario(scenario, expected_status):
    server.simulation_state["scenario"] = scenario

    snapshot = server.simulated_snapshot()
    upload, download = server.simulated_traffic_sample()

    assert snapshot["status"] == expected_status
    assert upload >= 0
    assert download >= 0


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("", ""),
        ("FORTINET, INC.", "Fortinet"),
        ("Huawei Technologies Co., Ltd.", "Huawei"),
        ("COMPAL ELECTRONICS, INC.", "Compal"),
        ("Cisco Systems", "Cisco"),
        ("Hewlett Packard Enterprise", "HP / Aruba"),
        ("TP-Link Corporation", "TP-Link"),
        ("Intel Corporate", "Intel"),
        ("Dell Inc.", "Dell"),
        ("Apple, Inc.", "Apple"),
        ("Samsung Electronics", "Samsung"),
        ("Realtek Semiconductor", "Realtek"),
        ("Synology Incorporated", "Synology"),
        ("QNAP Systems", "QNAP"),
        ("ACME DEVICES LTD.", "Acme Devices"),
    ],
)
def test_vendor_names_are_normalized_without_inventing_identity(raw, expected):
    assert server._clean_vendor_display(raw) == expected
